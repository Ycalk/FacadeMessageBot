"""Выгрузка почасовой статистики сообщений по пользователям в XLSX."""

from __future__ import annotations

import argparse
import asyncio
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy import Integer, cast, func, select, text

from db.models import Message, User
from db.session import async_session

_MSK = timedelta(hours=3)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"Некорректная дата '{value}'. Используйте формат YYYY-MM-DD."
        ) from exc


def _build_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Выгрузка статистики сообщений по пользователям в XLSX."
    )
    parser.add_argument(
        "--date-from",
        required=True,
        type=_parse_date,
        help="Дата начала диапазона (включительно), формат YYYY-MM-DD, МСК.",
    )
    parser.add_argument(
        "--date-to",
        required=True,
        type=_parse_date,
        help="Дата конца диапазона (включительно), формат YYYY-MM-DD, МСК.",
    )
    parser.add_argument(
        "--output",
        required=True,
        help="Путь к итоговому XLSX, например /home/app/data/user_hourly_stats.xlsx",
    )
    return parser.parse_args()


async def _load_rows(date_from: date, date_to: date) -> list:
    start_utc = datetime.combine(date_from, time.min) - _MSK
    end_utc = datetime.combine(date_to + timedelta(days=1), time.min) - _MSK

    created_at_msk = Message.created_at + text("interval '3 hours'")
    day_expr = func.date(created_at_msk)
    hour_expr = cast(func.extract("hour", created_at_msk), Integer)

    async with async_session() as session:
        result = await session.execute(
            select(
                day_expr.label("day_msk"),
                hour_expr.label("hour_msk"),
                Message.user_id,
                User.max_id,
                User.first_name,
                User.username,
                func.count(Message.id).label("messages_count"),
            )
            .join(User, User.id == Message.user_id)
            .where(Message.created_at >= start_utc, Message.created_at < end_utc)
            .group_by(
                day_expr,
                hour_expr,
                Message.user_id,
                User.max_id,
                User.first_name,
                User.username,
            )
            .order_by(day_expr, hour_expr, Message.user_id)
        )
        return result.all()


def _save_xlsx(rows: list, output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Почасовая статистика"

    headers = [
        "Дата (МСК)",
        "Час (МСК)",
        "user_id",
        "max_id",
        "first_name",
        "username",
        "Количество сообщений",
    ]
    widths = [14, 12, 10, 12, 20, 20, 22]

    for idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = width

    ws.freeze_panes = "A2"

    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=row.day_msk.strftime("%Y-%m-%d"))
        ws.cell(row=row_idx, column=2, value=f"{row.hour_msk:02d}:00")
        ws.cell(row=row_idx, column=3, value=row.user_id)
        ws.cell(row=row_idx, column=4, value=row.max_id)
        ws.cell(row=row_idx, column=5, value=row.first_name or "")
        ws.cell(row=row_idx, column=6, value=row.username or "")
        ws.cell(row=row_idx, column=7, value=row.messages_count)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


async def _main() -> None:
    args = _build_args()
    if args.date_to < args.date_from:
        raise SystemExit("--date-to не может быть раньше --date-from")

    rows = await _load_rows(args.date_from, args.date_to)
    output_path = Path(args.output)
    _save_xlsx(rows, output_path)
    print(f"OK: {len(rows)} строк(и) записано в {output_path}")


if __name__ == "__main__":
    asyncio.run(_main())
